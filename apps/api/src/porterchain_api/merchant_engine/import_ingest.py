"""Parse CSV / XLSX / XLS into tabular rows for route import."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from typing import Any


@dataclass
class TabularSheet:
    headers: list[str]
    rows: list[dict[str, Any]]
    sheet_name: str = "Sheet1"


def _norm_header(h: Any) -> str:
    return str(h or "").strip().lower().replace(" ", "_")


def parse_csv_bytes(data: bytes) -> TabularSheet:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return TabularSheet(headers=[], rows=[])
    headers = [_norm_header(h) for h in reader.fieldnames]
    rows: list[dict[str, Any]] = []
    for raw in reader:
        row = {_norm_header(k): (v.strip() if isinstance(v, str) else v) for k, v in raw.items() if k}
        if any(str(v or "").strip() for v in row.values()):
            rows.append(row)
    return TabularSheet(headers=headers, rows=rows)


def parse_xlsx_bytes(data: bytes) -> TabularSheet:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("excel_support_unavailable") from exc

    wb = load_workbook(filename=io.BytesIO(data), read_only=True, data_only=True)
    # densest sheet
    best_name = wb.sheetnames[0]
    best_rows = -1
    best: TabularSheet | None = None
    for name in wb.sheetnames:
        ws = wb[name]
        grid = list(ws.iter_rows(values_only=True))
        if not grid:
            continue
        headers = [_norm_header(c) for c in grid[0]]
        rows: list[dict[str, Any]] = []
        for line in grid[1:]:
            if line is None:
                continue
            row = {
                headers[i]: (str(line[i]).strip() if line[i] is not None else "")
                for i in range(min(len(headers), len(line)))
                if headers[i]
            }
            if any(str(v or "").strip() for v in row.values()):
                rows.append(row)
        if len(rows) > best_rows:
            best_rows = len(rows)
            best = TabularSheet(headers=[h for h in headers if h], rows=rows, sheet_name=name)
            best_name = name
    wb.close()
    return best or TabularSheet(headers=[], rows=[], sheet_name=best_name)


def parse_xls_bytes(data: bytes) -> TabularSheet:
    """Legacy .xls via xlrd when installed; otherwise clear error."""
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise ValueError("xls_support_unavailable") from exc

    book = xlrd.open_workbook(file_contents=data)
    best: TabularSheet | None = None
    best_rows = -1
    for idx in range(book.nsheets):
        sheet = book.sheet_by_index(idx)
        if sheet.nrows < 1:
            continue
        headers = [_norm_header(sheet.cell_value(0, c)) for c in range(sheet.ncols)]
        rows: list[dict[str, Any]] = []
        for r in range(1, sheet.nrows):
            row = {
                headers[c]: str(sheet.cell_value(r, c)).strip()
                for c in range(sheet.ncols)
                if headers[c]
            }
            if any(row.values()):
                rows.append(row)
        if len(rows) > best_rows:
            best_rows = len(rows)
            best = TabularSheet(headers=[h for h in headers if h], rows=rows, sheet_name=sheet.name)
    return best or TabularSheet(headers=[], rows=[])


def parse_upload(filename: str, data: bytes) -> TabularSheet:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return parse_csv_bytes(data)
    if name.endswith(".xlsx"):
        return parse_xlsx_bytes(data)
    if name.endswith(".xls"):
        return parse_xls_bytes(data)
    # sniff
    if data[:2] == b"PK":
        return parse_xlsx_bytes(data)
    return parse_csv_bytes(data)
