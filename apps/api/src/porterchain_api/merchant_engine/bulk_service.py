"""Bulk CSV/Excel import per MODULE_BREAKDOWN.md."""

from __future__ import annotations

import csv
import hashlib
import io
from datetime import UTC, datetime, timedelta
from typing import Any

from sqlalchemy.orm import Session

from porterchain_api.booking_engine._core import emit_event
from porterchain_api.config import Settings
from porterchain_api.domain.customer_goods import persist_vehicle_class
from porterchain_api.domain.merchant_states import BulkImportStatus
from porterchain_api.domain.states import OrderSource
from porterchain_api.merchant_engine import events as E
from porterchain_api.merchant_engine.booking_flow_service import MerchantBookingFlowService
from porterchain_api.merchant_engine.booking_service import MerchantBookingService
from porterchain_api.merchant_engine.import_geocode import geocode_stop
from porterchain_api.merchant_engine.rbac import MerchantContext
from porterchain_api.merchant_models import BulkImportJob
from porterchain_api.schemas_merchant import AddressInput, MerchantBookDeliveryRequest

REQUIRED_COLUMNS = {"pickup", "dropoff", "scheduled_at"}

# Re-upload protection: rows already imported by an earlier, confirmed job for the
# same merchant within this window are skipped (reported as "already_imported").
REUPLOAD_LOOKBACK_DAYS = 30
_PRIOR_JOB_SCAN_LIMIT = 200


def _norm(value: Any) -> str:
    return " ".join(str(value or "").split()).lower()


def row_fingerprint(row: dict[str, Any]) -> str:
    """Stable identity of a shipment row (case/whitespace-insensitive).

    pickup + dropoff + scheduled_at (+ internal_reference when present) — two rows
    for the same lane and time with different merchant references are distinct jobs.
    """
    parts = [_norm(row.get("pickup")), _norm(row.get("dropoff")), _norm(row.get("scheduled_at"))]
    ref = _norm(row.get("internal_reference"))
    if ref:
        parts.append(ref)
    return hashlib.sha256("|".join(parts).encode()).hexdigest()


class MerchantBulkService:
    def __init__(self) -> None:
        self._booking = MerchantBookingService()
        self._flow = MerchantBookingFlowService()

    def upload_file(
        self,
        db: Session,
        settings: Settings,
        ctx: MerchantContext,
        *,
        filename: str,
        content: bytes,
    ) -> BulkImportJob:
        rows = self._parse_rows(filename, content)
        return self._build_job(
            db,
            settings,
            ctx,
            filename=filename,
            rows=rows,
            file_sha256=hashlib.sha256(content).hexdigest(),
        )

    def upload_csv(
        self,
        db: Session,
        ctx: MerchantContext,
        *,
        filename: str,
        content: str,
    ) -> BulkImportJob:
        rows = list(csv.DictReader(io.StringIO(content)))
        normalized = [self._normalize_row(row) for row in rows]
        return self._build_job(
            db,
            None,
            ctx,
            filename=filename,
            rows=normalized,
            file_sha256=hashlib.sha256(content.encode()).hexdigest(),
        )

    def _prior_imports(
        self, db: Session, ctx: MerchantContext
    ) -> tuple[dict[str, str], dict[str, BulkImportJob]]:
        """Row fingerprints from recent confirmed jobs, and recent jobs by file hash."""
        since = datetime.now(UTC) - timedelta(days=REUPLOAD_LOOKBACK_DAYS)
        jobs = (
            db.query(BulkImportJob)
            .filter(
                BulkImportJob.merchant_id == ctx.merchant.id,
                BulkImportJob.kind == "classic",
                BulkImportJob.created_at >= since,
            )
            .order_by(BulkImportJob.created_at.desc())
            .limit(_PRIOR_JOB_SCAN_LIMIT)
            .all()
        )
        imported: dict[str, str] = {}
        by_file: dict[str, BulkImportJob] = {}
        for job in jobs:
            cfg = job.job_config if isinstance(job.job_config, dict) else {}
            file_hash = cfg.get("file_sha256")
            if file_hash and file_hash not in by_file:
                by_file[file_hash] = job
            if job.status != BulkImportStatus.CONFIRMED.value:
                continue
            for fp in cfg.get("row_fingerprints") or []:
                imported.setdefault(str(fp), job.id)
        return imported, by_file

    def _parse_rows(self, filename: str, content: bytes) -> list[dict[str, Any]]:
        if filename.lower().endswith(".xlsx"):
            return self._parse_xlsx(content)
        text = content.decode("utf-8", errors="replace")
        return [self._normalize_row(row) for row in csv.DictReader(io.StringIO(text))]

    def _parse_xlsx(self, content: bytes) -> list[dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("excel_support_unavailable") from exc

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            return []

        headers = [str(cell).strip().lower() if cell is not None else "" for cell in header_row]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if not any(values):
                continue
            row = {
                headers[i]: ("" if values[i] is None else str(values[i]).strip())
                for i in range(min(len(headers), len(values)))
                if headers[i]
            }
            rows.append(self._normalize_row(row))
        return rows

    def _normalize_row(self, row: dict[str, Any]) -> dict[str, Any]:
        return {str(k).strip().lower(): ("" if v is None else str(v).strip()) for k, v in row.items() if k}

    def _build_job(
        self,
        db: Session,
        settings: Settings | None,
        ctx: MerchantContext,
        *,
        filename: str,
        rows: list[dict[str, Any]],
        file_sha256: str | None = None,
    ) -> BulkImportJob:
        preview: list[dict[str, Any]] = []
        errors: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        seen_hashes: set[str] = set()
        duplicates = 0
        already_imported = 0
        imported, by_file = self._prior_imports(db, ctx)
        if file_sha256 and file_sha256 in by_file:
            prior = by_file[file_sha256]
            warnings.append(
                {
                    "code": "duplicate_file",
                    "prior_job_id": prior.id,
                    "prior_status": prior.status,
                    "prior_filename": prior.filename,
                    "prior_uploaded_at": prior.created_at.isoformat() if prior.created_at else None,
                    "message": (
                        "This exact file was already uploaded"
                        + (" and booked" if prior.status == BulkImportStatus.CONFIRMED.value else "")
                        + ". Rows that were already booked are skipped."
                    ),
                }
            )

        for i, row in enumerate(rows, start=2):
            missing = REQUIRED_COLUMNS - set(row.keys())
            if missing:
                errors.append({"row": i, "error": f"missing_columns:{','.join(sorted(missing))}"})
                continue

            row_hash = row_fingerprint(row)
            if row_hash in seen_hashes:
                duplicates += 1
                errors.append({"row": i, "error": "duplicate"})
                continue
            seen_hashes.add(row_hash)
            if row_hash in imported:
                duplicates += 1
                already_imported += 1
                errors.append({"row": i, "error": "already_imported", "prior_job_id": imported[row_hash]})
                continue

            pickup = self._address_from_row(row, prefix="pickup")
            dropoff = self._address_from_row(row, prefix="dropoff")
            address_errors = self._flow.validate_addresses(pickup, dropoff, require_coordinates=False)
            if address_errors:
                errors.append({"row": i, "error": "address_validation", "details": address_errors})
                continue

            recipient_id = row.get("recipient_id") or None
            if recipient_id:
                try:
                    self._flow.validate_recipient(db, ctx, recipient_id)
                except LookupError:
                    errors.append({"row": i, "error": "recipient_not_found", "recipient_id": recipient_id})
                    continue

            try:
                scheduled = datetime.fromisoformat(row["scheduled_at"].replace("Z", "+00:00"))
            except ValueError:
                errors.append({"row": i, "error": "invalid_scheduled_at"})
                continue

            body = MerchantBookDeliveryRequest(
                pickup=pickup,
                dropoff=dropoff,
                scheduled_at=scheduled,
                vehicle_class=persist_vehicle_class(row.get("vehicle_class")),
                package_type=row.get("package_type") or "looseParcel",
                weight_kg=float(row["weight_kg"]) if row.get("weight_kg") else None,
                internal_reference=row.get("internal_reference") or None,
                purchase_order_number=row.get("purchase_order_number") or None,
                cost_centre=row.get("cost_centre") or None,
                recipient_id=recipient_id,
            )

            row_preview: dict[str, Any] = {"row": i, **row, "_fingerprint": row_hash}
            if settings is not None:
                quote = self._flow.preview(db, settings, ctx, body)
                if not quote.get("valid"):
                    if quote.get("address_errors"):
                        errors.append(
                            {
                                "row": i,
                                "error": "address_not_found",
                                "message": "We could not find this address. Add the postal code or lat/lng columns.",
                                "details": quote["address_errors"],
                            }
                        )
                        continue
                    errors.append(
                        {
                            "row": i,
                            "error": quote.get("error") or "pricing_validation_failed",
                            "message": quote.get("message"),
                        }
                    )
                    continue
                row_preview["estimated_amount_cents"] = quote.get("amount_cents")
                row_preview["vehicle_class"] = quote.get("vehicle_class")
                row_preview["pricing_breakdown"] = quote.get("pricing_breakdown")

            preview.append(row_preview)

        if already_imported:
            warnings.append(
                {
                    "code": "rows_already_imported",
                    "count": already_imported,
                    "message": (
                        f"{already_imported} row(s) match shipments booked from an earlier upload "
                        f"in the last {REUPLOAD_LOOKBACK_DAYS} days and will not be booked again."
                    ),
                }
            )

        job = BulkImportJob(
            merchant_id=ctx.merchant.id,
            status=BulkImportStatus.PREVIEW.value,
            kind="classic",
            filename=filename,
            total_rows=len(rows),
            valid_rows=len(preview),
            error_rows=len(errors) - duplicates,
            duplicate_rows=duplicates,
            preview=preview[:50],
            errors=errors[:100],
            job_config={
                "file_sha256": file_sha256,
                "row_fingerprints": [r["_fingerprint"] for r in preview[:50]],
                "warnings": warnings,
            },
        )
        db.add(job)
        db.commit()
        db.refresh(job)
        return job

    def _address_from_row(self, row: dict[str, Any], *, prefix: str) -> AddressInput:
        """CSV rows usually carry text addresses only; geocode them so the row can be priced."""
        lat_raw = row.get(f"{prefix}_lat")
        lng_raw = row.get(f"{prefix}_lng")
        lat = float(lat_raw) if lat_raw else None
        lng = float(lng_raw) if lng_raw else None
        if lat is None or lng is None:
            geo = geocode_stop(address=row[prefix], source="csv")
            if geo.lat is not None and geo.lng is not None:
                return AddressInput(
                    formatted=geo.formatted or row[prefix], lat=geo.lat, lng=geo.lng, postal=geo.postal
                )
        return AddressInput(formatted=row[prefix], lat=lat, lng=lng)

    def confirm_bulk(
        self,
        db: Session,
        settings: Settings,
        ctx: MerchantContext,
        job_id: str,
        *,
        is_sandbox: bool = False,
    ) -> BulkImportJob:
        job = (
            db.query(BulkImportJob)
            .filter(BulkImportJob.id == job_id, BulkImportJob.merchant_id == ctx.merchant.id)
            .with_for_update()
            .first()
        )
        if not job:
            raise LookupError("bulk_job_not_found")
        if job.status == BulkImportStatus.CONFIRMED.value:
            # Idempotent: a double-click / retried confirm never books twice.
            return job

        # Another upload of the same rows may have been confirmed since this preview.
        imported, _ = self._prior_imports(db, ctx)
        order_ids: list[str] = []
        confirm_errors: list[dict[str, Any]] = []
        for row in job.preview:
            fp = row.get("_fingerprint") or row_fingerprint(row)
            if fp in imported and imported[fp] != job.id:
                confirm_errors.append(
                    {"row": row.get("row"), "error": "already_imported", "prior_job_id": imported[fp]}
                )
                continue
            try:
                scheduled = datetime.fromisoformat(row["scheduled_at"].replace("Z", "+00:00"))
            except ValueError:
                scheduled = datetime.now(UTC)
            body = MerchantBookDeliveryRequest(
                pickup=self._address_from_row(row, prefix="pickup"),
                dropoff=self._address_from_row(row, prefix="dropoff"),
                scheduled_at=scheduled,
                vehicle_class=persist_vehicle_class(row.get("vehicle_class")),
                package_type=row.get("package_type") or "looseParcel",
                weight_kg=float(row["weight_kg"]) if row.get("weight_kg") else None,
                internal_reference=row.get("internal_reference"),
                purchase_order_number=row.get("purchase_order_number"),
                cost_centre=row.get("cost_centre"),
                recipient_id=row.get("recipient_id"),
                is_sandbox=is_sandbox,
            )
            preview = self._flow.preview(db, settings, ctx, body)
            if not preview.get("valid"):
                confirm_errors.append({"row": row.get("row"), "error": preview.get("error")})
                continue
            order = self._booking.create_shipment(
                db,
                settings,
                ctx,
                body,
                order_source=OrderSource.CSV.value,
                sandbox=is_sandbox,
            )
            order_ids.append(order.id)

        job.status = BulkImportStatus.CONFIRMED.value
        job.order_ids = order_ids
        if confirm_errors:
            job.errors = list(job.errors or []) + confirm_errors
        db.commit()

        emit_event(
            db,
            event_type=E.MERCHANT_BULK_BOOKING_CREATED,
            aggregate_type="bulk_import",
            aggregate_id=job.id,
            correlation_id=ctx.merchant.id,
            actor_type="merchant",
            actor_id=ctx.user.id,
            payload={"order_count": len(order_ids)},
        )
        db.commit()
        db.refresh(job)
        return job

    @staticmethod
    def upload_payload(job: BulkImportJob) -> dict[str, Any]:
        return {
            "job_id": job.id,
            "status": job.status,
            "total_rows": job.total_rows,
            "valid_rows": job.valid_rows,
            "error_rows": job.error_rows,
            "duplicate_rows": job.duplicate_rows,
            "preview": [{k: v for k, v in row.items() if k != "_fingerprint"} for row in job.preview or []],
            "errors": job.errors,
            "warnings": list((job.job_config or {}).get("warnings") or [])
            if isinstance(job.job_config, dict)
            else [],
        }
